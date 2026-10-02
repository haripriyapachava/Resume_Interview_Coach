import os
import time
import tempfile
from pathlib import Path

from flask import Flask, request, render_template_string
from google import genai
from pypdf import PdfReader
from docx import Document
from werkzeug.utils import secure_filename
import markdown

app = Flask(__name__)

# ============================================================
# CONFIGURATION
# ============================================================

MODEL = "gemini-3.5-flash-lite"

ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


# ============================================================
# GEMINI CLIENT
# ============================================================

API_KEY = os.getenv("GEMINI_API_KEY")

if API_KEY:
    client = genai.Client(api_key=API_KEY)
else:
    client = None


# ============================================================
# HTML TEMPLATE
# ============================================================

HTML = """
<!DOCTYPE html>
<html lang="en">

<head>

<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">

<title>Resume Interview Coach</title>

<style>

* {
    box-sizing: border-box;
}

body {
    margin: 0;
    font-family: Inter, Arial, sans-serif;
    background: #f5f7fb;
    color: #1f2937;
}

/* ----------------------------------------------------------
   HEADER
---------------------------------------------------------- */

.header {
    background: white;
    border-bottom: 1px solid #e5e7eb;
    padding: 22px 8%;
}

.logo {
    font-size: 22px;
    font-weight: 700;
    color: #111827;
}

.logo span {
    color: #6366f1;
}

/* ----------------------------------------------------------
   MAIN
---------------------------------------------------------- */

.container {
    width: 90%;
    max-width: 1050px;
    margin: 45px auto;
}

.hero {
    text-align: center;
    margin-bottom: 35px;
}

.hero h1 {
    font-size: 38px;
    margin-bottom: 12px;
    color: #111827;
}

.hero p {
    color: #6b7280;
    font-size: 17px;
}

/* ----------------------------------------------------------
   UPLOAD CARD
---------------------------------------------------------- */

.upload-card {
    background: white;
    border-radius: 18px;
    padding: 35px;
    border: 1px solid #e5e7eb;
    box-shadow: 0 10px 30px rgba(0,0,0,0.05);
    margin-bottom: 35px;
}

.upload-area {
    border: 2px dashed #c7d2fe;
    background: #fafaff;
    border-radius: 15px;
    padding: 35px;
    text-align: center;
    transition: 0.2s;
}

.upload-area:hover {
    border-color: #6366f1;
    background: #f8f8ff;
}

.upload-icon {
    font-size: 42px;
    margin-bottom: 10px;
}

.upload-area h3 {
    margin: 8px 0;
    color: #111827;
}

.upload-area p {
    color: #6b7280;
    font-size: 14px;
}

input[type="file"] {
    margin: 18px 0;
    padding: 10px;
    width: 100%;
    max-width: 420px;
}

.button {
    border: none;
    background: #6366f1;
    color: white;
    padding: 14px 26px;
    border-radius: 10px;
    font-size: 15px;
    font-weight: 600;
    cursor: pointer;
    transition: 0.2s;
}

.button:hover {
    background: #4f46e5;
    transform: translateY(-1px);
}

/* ----------------------------------------------------------
   RESULT AREA
---------------------------------------------------------- */

.result-container {
    margin-top: 30px;
}

.result-title {
    margin-bottom: 18px;
}

.result-title h2 {
    margin: 0;
    font-size: 25px;
}

.result-title p {
    color: #6b7280;
    margin-top: 7px;
}

/* ----------------------------------------------------------
   AI OUTPUT
---------------------------------------------------------- */

.answer-box {
    background: white;
    border: 1px solid #e5e7eb;
    border-radius: 16px;
    padding: 30px;
    box-shadow: 0 8px 25px rgba(0,0,0,0.04);
}

/* Markdown rendered content */

.answer-box h1 {
    font-size: 27px;
    color: #111827;
    margin-top: 10px;
    margin-bottom: 18px;
    border-bottom: 1px solid #e5e7eb;
    padding-bottom: 12px;
}

.answer-box h2 {
    font-size: 22px;
    color: #4338ca;
    margin-top: 32px;
    margin-bottom: 15px;
    padding-bottom: 8px;
    border-bottom: 1px solid #eef0f5;
}

.answer-box h3 {
    font-size: 18px;
    color: #374151;
    margin-top: 25px;
    margin-bottom: 12px;
}

.answer-box p {
    line-height: 1.7;
    color: #374151;
    margin: 9px 0;
}

.answer-box ul {
    padding-left: 24px;
    margin-top: 8px;
}

.answer-box li {
    margin: 9px 0;
    line-height: 1.6;
}

.answer-box strong {
    color: #111827;
}

.answer-box hr {
    border: none;
    border-top: 1px solid #e5e7eb;
    margin: 25px 0;
}

/* Question styling */

.answer-box li strong:first-child {
    color: #4f46e5;
}

/* Code */

.answer-box code {
    background: #f3f4f6;
    padding: 3px 6px;
    border-radius: 5px;
    font-size: 14px;
}

.answer-box pre {
    background: #111827;
    color: white;
    padding: 18px;
    border-radius: 10px;
    overflow-x: auto;
}

/* ----------------------------------------------------------
   INFO CARDS
---------------------------------------------------------- */

.info-row {
    display: flex;
    gap: 15px;
    margin-top: 25px;
}

.info-card {
    flex: 1;
    background: #fafafa;
    border: 1px solid #e5e7eb;
    border-radius: 12px;
    padding: 16px;
}

.info-card strong {
    display: block;
    margin-bottom: 5px;
}

.info-card span {
    color: #6b7280;
    font-size: 13px;
}

/* ----------------------------------------------------------
   ERROR
---------------------------------------------------------- */

.error {
    background: #fff1f2;
    border: 1px solid #fecdd3;
    color: #be123c;
    padding: 18px;
    border-radius: 12px;
    margin-top: 20px;
}

/* ----------------------------------------------------------
   LOADING
---------------------------------------------------------- */

.loading {
    text-align: center;
    padding: 30px;
    color: #6366f1;
    font-weight: 600;
}

/* ----------------------------------------------------------
   FOOTER
---------------------------------------------------------- */

.footer {
    text-align: center;
    color: #9ca3af;
    font-size: 13px;
    margin: 45px 0 25px;
}

/* ----------------------------------------------------------
   MOBILE
---------------------------------------------------------- */

@media(max-width: 700px) {

    .hero h1 {
        font-size: 30px;
    }

    .upload-card {
        padding: 20px;
    }

    .answer-box {
        padding: 20px;
    }

    .info-row {
        flex-direction: column;
    }
}

</style>

</head>


<body>

<div class="header">

    <div class="logo">
        Resume <span>Interview Coach</span>
    </div>

</div>


<div class="container">

    <div class="hero">

        <h1>Prepare for Your Interview</h1>

        <p>
            Upload your resume and get interview questions,
            follow-ups, and interview-ready answers based on your resume.
        </p>

    </div>


    <div class="upload-card">

        <form method="POST"
              enctype="multipart/form-data">

            <div class="upload-area">

                <div class="upload-icon">
                    📄
                </div>

                <h3>
                    Upload Your Resume
                </h3>

                <p>
                    PDF, DOCX or TXT • Maximum 10 MB
                </p>

                <input
                    type="file"
                    name="resume"
                    accept=".pdf,.docx,.txt"
                    required
                >

                <br>

                <button
                    type="submit"
                    class="button">

                    🚀 Prepare Me for My Interview

                </button>

            </div>

        </form>

    </div>


    {% if error %}

        <div class="error">
            ⚠️ {{ error }}
        </div>

    {% endif %}


    {% if result %}

    <div class="result-container">

        <div class="result-title">

            <h2>
                🎯 Your Interview Preparation
            </h2>

            <p>
                Generated from the information available in your resume.
            </p>

        </div>


        <div class="answer-box">

            {{ result|safe }}

        </div>

    </div>

    {% endif %}


    <div class="info-row">

        <div class="info-card">

            <strong>📌 Resume Based</strong>

            <span>
                Questions are connected to your resume.
            </span>

        </div>


        <div class="info-card">

            <strong>💻 Technical</strong>

            <span>
                Skills and technologies are covered.
            </span>

        </div>


        <div class="info-card">

            <strong>🔁 Follow-ups</strong>

            <span>
                Practice deeper interview questions.
            </span>

        </div>

    </div>


    <div class="footer">

        Resume Interview Coach • AI-powered interview preparation

    </div>

</div>

</body>

</html>
"""


# ============================================================
# FILE HELPERS
# ============================================================

def allowed_file(filename):

    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower()
        in ALLOWED_EXTENSIONS
    )


# ============================================================
# RESUME TEXT EXTRACTION
# ============================================================

def extract_resume_text(filepath, extension):

    text = ""

    # ---------------- PDF ----------------

    if extension == "pdf":

        reader = PdfReader(filepath)

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:
                text += page_text + "\n"


    # ---------------- DOCX ----------------

    elif extension == "docx":

        document = Document(filepath)

        for paragraph in document.paragraphs:

            if paragraph.text.strip():
                text += paragraph.text + "\n"

        # Also read tables

        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    if cell.text.strip():
                        row_text.append(cell.text.strip())

                if row_text:
                    text += " | ".join(row_text) + "\n"


    # ---------------- TXT ----------------

    elif extension == "txt":

        with open(
            filepath,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:

            text = file.read()


    return text.strip()


# ============================================================
# GEMINI ANALYSIS
# ============================================================

def analyze_resume(resume_text):

    if not client:

        return """
        ⚠️ Gemini API key is not configured.

        Please set the GEMINI_API_KEY environment variable before running
        the application.
        """

    prompt = f"""
You are an expert technical interviewer and interview coach.

Your job is to prepare a candidate for an interview using ONLY the
information contained in the resume below.

IMPORTANT RULES:

1. Do NOT invent projects, technologies, companies, internships,
   certifications, achievements, responsibilities, numbers, or experience.

2. If something is unclear from the resume, explicitly say:
   "This is not clear from the resume."

3. For generic technical concepts, you may provide accurate
   interview-ready explanations.

4. Every answer should be concise enough to speak naturally in an
   interview.

5. The candidate should be able to understand the answer easily.

6. Do not give a generic interview guide.

7. Make the preparation specifically connected to THIS resume.

8. Generate BOTH QUESTIONS AND ANSWERS.

9. Use clean Markdown headings and bullet points.
   Do NOT use giant ASCII separators such as:
   =========================
   -------------------------
   Do NOT number every section excessively.

RESUME:

{resume_text}


Prepare the following:

# 1. Resume Breakdown

Briefly identify:

- Education
- Technical skills
- Programming languages
- Frameworks
- Databases
- Projects
- Internships / experience
- Certifications
- Achievements
- Tools
- Other important resume items

Only mention things actually present in the resume.


# 2. Questions From My Resume

For every important resume item, generate likely interviewer questions.

For each question use this format:

**Q1. Question**

**Answer:** Interview-ready answer based on the resume.

Include questions such as:

- Tell me about this.
- Why did you choose this?
- What exactly did you do?
- What was your role?
- What did you learn?
- What challenges did you face?
- What would you improve?

Do NOT ask about things that are not present in the resume.


# 3. Technical Questions

For EVERY major technical skill or technology listed in the resume:

Create questions at:

- Basic level
- Intermediate level
- Deeper interview level

For each:

**Question:** ...

**Answer:** ...

Keep answers simple and interview-ready.

If a technology is listed but the resume does not provide enough personal detail,
ask concept-based questions rather than inventing personal experience.


# 4. Project Interview Questions

For EACH project listed in the resume, prepare:

- Project explanation
- Problem statement
- Why the project was built
- Candidate's role
- Technologies used
- Architecture / workflow
- Database
- APIs
- Important implementation decisions
- Challenges
- How challenges were solved
- Testing
- Limitations
- Future improvements

For every item give:

**Question:** ...

**Answer:** ...


# 5. Project Follow-Up Chains

For each major project create realistic interviewer follow-ups.

Example:

**Main Question:** Explain your project.

**Answer:** ...

**Follow-up:** Why did you choose this technology?

**Answer:** ...

**Deeper Follow-up:** What would happen if the number of users increased significantly?

**Answer:** ...

Only use project facts supported by the resume.


# 6. Resume Skill Cross-Questions

Create questions where an interviewer connects two or more skills
from the resume.

Example:

If the resume contains Python + Flask + SQL:

**Question:** How did Python, Flask and SQL work together in your project?

**Answer:** ...

Only create combinations supported by the resume.


# 7. HR Questions Based on My Resume

Prepare questions such as:

- Tell me about yourself.
- Why this role?
- Tell me about your project.
- What was your biggest learning?
- Tell me about a challenge.
- What are your strengths?
- What are you currently improving?

Connect the questions to the actual resume where possible.

Provide interview-ready answers.


# 8. Questions I Must Be Ready to Defend

Identify important claims, technologies, projects, achievements,
or responsibilities on the resume that an interviewer may ask the
candidate to explain in detail.

For each:

**What you should be ready to explain:** ...

**Possible Question:** ...

**Answer:** ...


# 9. Top Questions to Practice First

Give a practical list of the most important questions to practice first.

Do NOT rank the candidate.

This is only a practice order.

For example:

1. Explain your main project.
2. Explain your role in the project.
3. Explain your strongest technical skill.
4. Explain an important challenge you solved.

Use questions relevant to THIS resume.


# 10. Mock Interview

Finish with ONE realistic first interview question based directly
on the resume.

Do not provide the answer for this final mock question.

The candidate should answer it themselves.


Keep the entire response organized, readable, and interview-focused.
"""


    # ========================================================
    # RETRY LOGIC
    # ========================================================

    for attempt in range(3):

        try:

            response = client.models.generate_content(
                model=MODEL,
                contents=prompt
            )

            if response and response.text:

                return response.text

            return "Gemini returned an empty response."

        except Exception as e:

            error_text = str(e)

            if (
                "503" in error_text
                or "UNAVAILABLE" in error_text
            ):

                if attempt < 2:

                    time.sleep(5)

                    continue

                return (
                    "⚠️ Gemini is temporarily experiencing high demand. "
                    "Please wait a little and try again."
                )

            return f"⚠️ Gemini error: {error_text}"

    return "⚠️ Unable to generate interview preparation."


# ============================================================
# MAIN ROUTE
# ============================================================

@app.route("/", methods=["GET", "POST"])
def home():

    result = None
    error = None

    if request.method == "POST":

        # Check file

        if "resume" not in request.files:

            error = "Please upload your resume."

            return render_template_string(
                HTML,
                result=result,
                error=error
            )


        file = request.files["resume"]


        if file.filename == "":

            error = "Please select a resume file."

            return render_template_string(
                HTML,
                result=result,
                error=error
            )


        if not allowed_file(file.filename):

            error = (
                "Unsupported file type. "
                "Please upload PDF, DOCX or TXT."
            )

            return render_template_string(
                HTML,
                result=result,
                error=error
            )


        # ----------------------------------------------------
        # Save temporarily
        # ----------------------------------------------------

        filename = secure_filename(file.filename)

        extension = filename.rsplit(
            ".",
            1
        )[1].lower()


        temp_dir = tempfile.gettempdir()

        filepath = os.path.join(
            temp_dir,
            filename
        )


        try:

            file.save(filepath)


            # ------------------------------------------------
            # Extract resume
            # ------------------------------------------------

            resume_text = extract_resume_text(
                filepath,
                extension
            )


            if not resume_text:

                error = (
                    "I could not extract readable text from this resume. "
                    "Please try another file."
                )

                return render_template_string(
                    HTML,
                    result=result,
                    error=error
                )


            # ------------------------------------------------
            # Analyze
            # ------------------------------------------------

            raw_result = analyze_resume(
                resume_text
            )


            # ------------------------------------------------
            # Convert Markdown → HTML
            # ------------------------------------------------

            result = markdown.markdown(
                raw_result,
                extensions=[
                    "extra",
                    "sane_lists"
                ]
            )


        except Exception as e:

            error = f"Something went wrong: {str(e)}"


        finally:

            # ------------------------------------------------
            # Delete temporary file
            # ------------------------------------------------

            try:

                if os.path.exists(filepath):

                    os.remove(filepath)

            except Exception:

                pass


    return render_template_string(
        HTML,
        result=result,
        error=error
    )


# ============================================================
# RUN APPLICATION
# ============================================================

if __name__ == "__main__":

    port = int(
        os.environ.get(
            "PORT",
            5000
        )
    )

    app.run(
        host="0.0.0.0",
        port=port
    )